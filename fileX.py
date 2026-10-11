#!/usr/bin/env python3
"""
FileX - Professional File Converter & Compressor
================================================

Install the requirements once:

    pip install pillow reportlab pandas openpyxl pymupdf

Run:

    python main.py

Features
--------
* Convert  : JPG, PNG, WEBP, BMP, GIF, TIFF, ICO  <->  each other, + PDF
             PDF -> TXT / PNG / JPG,  TXT -> PDF
             CSV <-> XLSX <-> JSON
* Compress : images (quality + resize), PDF optimisation, ZIP archives
* History  : every operation is logged, open files / folders directly
* Settings : default output folder, ask-where-to-save, default quality
"""

from __future__ import annotations

import importlib.util
import json
import os
import queue
import subprocess
import sys
import threading
import zipfile
from contextlib import contextmanager
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from tkinter import font as tkfont

# =============================================================================
# 1.  CONSTANTS & THEME
# =============================================================================

APP_NAME = "FileX"
APP_VERSION = "3.0"

# Dark charcoal + burnt orange palette (no white / blue anywhere)
BG = "#0E0E10"
SIDEBAR = "#121215"
CARD = "#18181C"
CARD_ALT = "#1F1F24"
HOVER = "#2A2A31"
BORDER = "#2A2A31"
NAV_ACTIVE = "#1D1814"
TEXT = "#ECE7DF"
MUTED = "#8E8A94"
DIM = "#5E5B64"
ACCENT = "#E8590C"          # dark orange
ACCENT_HOVER = "#FF7424"
ACCENT_DIM = "#3A2012"
ON_ACCENT = "#170C05"
SUCCESS = "#3DBE8B"
ERROR = "#F0524F"
WARNING = "#F5A524"

FONT_FAMILY = "Segoe UI"


def F(size: int = 10, weight: str = "normal"):
    return (FONT_FAMILY, size, weight)


DATA_DIR = Path.home() / ".filex"
SETTINGS_FILE = DATA_DIR / "settings.json"
HISTORY_FILE = DATA_DIR / "history.json"

# name -> (extension, category, description)
FORMATS = {
    "JPG": (".jpg", "Images", "JPEG photo format"),
    "PNG": (".png", "Images", "Lossless, supports transparency"),
    "WEBP": (".webp", "Images", "Modern compact web image"),
    "BMP": (".bmp", "Images", "Uncompressed bitmap"),
    "GIF": (".gif", "Images", "256-colour image"),
    "TIFF": (".tiff", "Images", "High-quality print image"),
    "ICO": (".ico", "Images", "Application / website icon"),
    "PDF": (".pdf", "Documents", "Portable document"),
    "TXT": (".txt", "Documents", "Plain text"),
    "CSV": (".csv", "Data", "Comma-separated table"),
    "XLSX": (".xlsx", "Data", "Excel workbook"),
    "JSON": (".json", "Data", "Structured data"),
}
EXT_ALIASES = {".jpeg": "JPG", ".jpe": "JPG", ".tif": "TIFF", ".xls": "XLSX"}
IMAGE_FORMATS = [n for n, v in FORMATS.items() if v[1] == "Images"]
INPUT_EXTENSIONS = sorted({v[0] for v in FORMATS.values()} | set(EXT_ALIASES))

REQUIREMENTS = {
    "PIL": "pillow",
    "reportlab": "reportlab",
    "pandas": "pandas",
    "openpyxl": "openpyxl",
}

DEFAULT_SETTINGS = {
    "output_folder": str(Path.home() / "Documents" / "FileX"),
    "ask_location": False,
    "quality": 85,
}


# =============================================================================
# 2.  HELPERS
# =============================================================================

def format_of(ext: str):
    """Return the canonical format name for a file extension (or None)."""
    ext = ext.lower()
    if ext in EXT_ALIASES:
        return EXT_ALIASES[ext]
    for name, (e, _, _) in FORMATS.items():
        if e == ext:
            return name
    return None


def available_targets(ext: str) -> list:
    """Formats a file with this extension can be converted to."""
    fmt = format_of(ext)
    if fmt is None:
        return []
    if fmt in IMAGE_FORMATS:
        return [f for f in IMAGE_FORMATS if f != fmt] + ["PDF"]
    return {
        "PDF": ["TXT", "PNG", "JPG"],
        "TXT": ["PDF"],
        "CSV": ["XLSX", "JSON"],
        "XLSX": ["CSV", "JSON"],
        "JSON": ["CSV", "XLSX"],
    }.get(fmt, [])


def human_size(num: float) -> str:
    size = float(num)
    for unit in ("B", "KB", "MB", "GB"):
        if size < 1024:
            return f"{int(size)} B" if unit == "B" else f"{size:.1f} {unit}"
        size /= 1024
    return f"{size:.1f} TB"


def size_change(before: int, after: int) -> str:
    if before <= 0:
        return ""
    pct = (after - before) / before * 100
    return f"{pct:+.1f}%".replace("-", "\u2212")


def shorten(text: str, limit: int = 34) -> str:
    if len(text) <= limit:
        return text
    keep = limit - 1
    head = keep // 2 + keep % 2
    return text[:head] + "\u2026" + text[-(keep - head):]


def unique_path(path: Path) -> Path:
    """Return a path that does not exist yet: name.ext, name (1).ext ..."""
    if not path.exists():
        return path
    n = 1
    while True:
        candidate = path.with_name(f"{path.stem} ({n}){path.suffix}")
        if not candidate.exists():
            return candidate
        n += 1


def open_path(path) -> None:
    path = str(path)
    if sys.platform == "win32":
        os.startfile(path)  # type: ignore[attr-defined]
    elif sys.platform == "darwin":
        subprocess.Popen(["open", path])
    else:
        subprocess.Popen(["xdg-open", path])


def reveal_path(path) -> None:
    path = Path(path)
    if sys.platform == "win32":
        subprocess.Popen(f'explorer /select,"{path}"')
    elif sys.platform == "darwin":
        subprocess.Popen(["open", "-R", str(path)])
    else:
        subprocess.Popen(["xdg-open", str(path.parent)])


def missing_dependencies() -> list:
    missing = [pip for mod, pip in REQUIREMENTS.items()
               if importlib.util.find_spec(mod) is None]
    if not (importlib.util.find_spec("pymupdf") or importlib.util.find_spec("fitz")):
        missing.append("pymupdf")
    return missing


def friendly_error(exc: BaseException) -> str:
    """Turn technical exceptions into messages a normal user understands."""
    name = type(exc).__name__
    if isinstance(exc, ImportError):
        mod = getattr(exc, "name", None) or ""
        pip = REQUIREMENTS.get(mod, mod or "the required package")
        return f"A required library is missing.\nInstall it with:  pip install {pip}"
    if isinstance(exc, PermissionError):
        return ("Permission denied. The file may be open in another program, "
                "or the folder is write-protected.")
    if isinstance(exc, FileNotFoundError):
        return "The file or folder could not be found."
    if name == "UnidentifiedImageError":
        return "This image is damaged or is not a valid image file."
    if name == "EmptyDataError":
        return "The file contains no data."
    if name in ("JSONDecodeError",):
        return "This JSON file is not valid JSON."
    return str(exc) or name


@contextmanager
def atomic_output(dst: Path):
    """Write to a temp file first, then move it into place.

    A failed or interrupted conversion can therefore never leave a broken file
    behind, and an existing file is only replaced once the new one is complete.
    """
    tmp = dst.with_name(dst.name + ".part")
    try:
        yield tmp
        os.replace(tmp, dst)
    finally:
        if tmp.exists():
            try:
                tmp.unlink()
            except OSError:
                pass


# =============================================================================
# 3.  STORAGE (settings + history)
# =============================================================================

def _read_json(path: Path, default):
    try:
        with open(path, "r", encoding="utf-8") as fh:
            return json.load(fh)
    except (OSError, ValueError):
        return default


def _write_json(path: Path, data) -> None:
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        tmp = path.with_suffix(".tmp")
        with open(tmp, "w", encoding="utf-8") as fh:
            json.dump(data, fh, indent=2, ensure_ascii=False)
        os.replace(tmp, path)
    except OSError as exc:
        print(f"Could not save {path.name}: {exc}")


class Settings(dict):
    def __init__(self):
        super().__init__(DEFAULT_SETTINGS)
        stored = _read_json(SETTINGS_FILE, {})
        if isinstance(stored, dict):
            for key in DEFAULT_SETTINGS:
                if key in stored:
                    self[key] = stored[key]

    def save(self):
        _write_json(SETTINGS_FILE, dict(self))


class History:
    LIMIT = 100

    def __init__(self):
        data = _read_json(HISTORY_FILE, [])
        self.items = data if isinstance(data, list) else []

    def add(self, operation: str, src: Path, dst: Path):
        self.items.insert(0, {
            "operation": operation,
            "input": src.name,
            "output": str(dst.resolve()),
            "time": datetime.now().strftime("%d %b %Y, %I:%M %p"),
        })
        del self.items[self.LIMIT:]
        _write_json(HISTORY_FILE, self.items)

    def clear(self):
        self.items = []
        _write_json(HISTORY_FILE, self.items)


# =============================================================================
# 4.  ENGINE  (pure functions, no GUI code)
# =============================================================================

@dataclass
class Result:
    outputs: list = field(default_factory=list)
    note: str = ""

    @property
    def primary(self) -> Path:
        return self.outputs[0]

    @property
    def total_size(self) -> int:
        return sum(p.stat().st_size for p in self.outputs)


def _pymupdf():
    try:
        import pymupdf
    except ImportError:
        import fitz as pymupdf  # older PyMuPDF versions
    return pymupdf


def _flatten(im):
    """Remove transparency (JPG/BMP/PDF cannot store it) using a plain background."""
    from PIL import Image
    if im.mode in ("RGBA", "LA") or (im.mode == "P" and "transparency" in im.info):
        rgba = im.convert("RGBA")
        base = Image.new("RGB", rgba.size, (255, 255, 255))
        base.paste(rgba, mask=rgba.getchannel("A"))
        return base
    return im.convert("RGB")


def convert_image(src: Path, dst: Path, target: str, quality: int) -> Result:
    from PIL import Image, ImageOps

    pil_names = {"JPG": "JPEG", "PNG": "PNG", "WEBP": "WEBP", "BMP": "BMP",
                 "GIF": "GIF", "TIFF": "TIFF", "ICO": "ICO", "PDF": "PDF"}
    with Image.open(src) as opened:
        im = ImageOps.exif_transpose(opened)
        if im.mode not in ("RGB", "RGBA", "L", "LA", "P", "1"):
            im = im.convert("RGBA")

        with atomic_output(dst) as tmp:
            if target == "JPG":
                _flatten(im).save(tmp, format="JPEG", quality=quality, optimize=True)
            elif target == "BMP":
                _flatten(im).save(tmp, format="BMP")
            elif target == "PDF":
                _flatten(im).save(tmp, format="PDF", resolution=100.0)
            elif target == "WEBP":
                im.save(tmp, format="WEBP", quality=quality)
            elif target == "PNG":
                im.save(tmp, format="PNG", optimize=True)
            elif target == "TIFF":
                im.convert("RGBA" if "A" in im.mode else "RGB").save(
                    tmp, format="TIFF", compression="tiff_lzw")
            elif target == "GIF":
                im.save(tmp, format="GIF")
            elif target == "ICO":
                rgba = im.convert("RGBA")
                side = max(rgba.size)
                canvas_img = Image.new("RGBA", (side, side), (0, 0, 0, 0))
                canvas_img.paste(rgba, ((side - rgba.width) // 2, (side - rgba.height) // 2))
                canvas_img = canvas_img.resize((256, 256), Image.LANCZOS)
                canvas_img.save(tmp, format="ICO",
                                sizes=[(16, 16), (32, 32), (48, 48), (64, 64),
                                       (128, 128), (256, 256)])
            else:
                raise ValueError(f"Unsupported image format: {target}")
    return Result([dst])


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return raw.decode(enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return raw.decode("latin-1", errors="replace")


_FONT_CANDIDATES = [
    r"C:\Windows\Fonts\segoeui.ttf",
    r"C:\Windows\Fonts\arial.ttf",
    "/Library/Fonts/Arial.ttf",
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
]


def _wrap(line: str, font: str, size: float, max_width: float, width_of) -> list:
    if not line.strip():
        return [""]
    out, current = [], ""
    for word in line.split(" "):
        trial = f"{current} {word}" if current else word
        if width_of(trial, font, size) <= max_width:
            current = trial
            continue
        if current:
            out.append(current)
        while width_of(word, font, size) > max_width:      # very long word
            k = len(word)
            while k > 1 and width_of(word[:k], font, size) > max_width:
                k -= 1
            out.append(word[:k])
            word = word[k:]
        current = word
    out.append(current)
    return out


def txt_to_pdf(src: Path, dst: Path) -> Result:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    text = _read_text(src)
    font_name = "Helvetica"
    for candidate in _FONT_CANDIDATES:
        if os.path.isfile(candidate):
            try:
                pdfmetrics.registerFont(TTFont("FileXText", candidate))
                font_name = "FileXText"
                break
            except Exception:
                continue

    size, leading, margin = 11, 16, 56
    page_w, page_h = A4
    max_w = page_w - 2 * margin
    width_of = pdfmetrics.stringWidth

    def footer(pdf, number):
        pdf.setFont(font_name, 8)
        pdf.setFillGray(0.5)
        pdf.drawCentredString(page_w / 2, 28, str(number))
        pdf.setFillGray(0)
        pdf.setFont(font_name, size)

    with atomic_output(dst) as tmp:
        pdf = canvas.Canvas(str(tmp), pagesize=A4)
        pdf.setTitle(src.stem)
        pdf.setFont(font_name, size)
        y, page = page_h - margin, 1
        for raw_line in text.splitlines():
            line = raw_line.expandtabs(4).rstrip()
            if font_name == "Helvetica":
                line = line.encode("cp1252", "replace").decode("cp1252")
            for piece in _wrap(line, font_name, size, max_w, width_of):
                if y < margin:
                    footer(pdf, page)
                    pdf.showPage()
                    page += 1
                    pdf.setFont(font_name, size)
                    y = page_h - margin
                pdf.drawString(margin, y, piece)
                y -= leading
        footer(pdf, page)
        pdf.save()
    return Result([dst], note=f"{page} page{'s' if page != 1 else ''}")


def pdf_to_txt(src: Path, dst: Path) -> Result:
    pymupdf = _pymupdf()
    doc = pymupdf.open(str(src))
    try:
        if doc.is_encrypted:
            raise ValueError("This PDF is password-protected. Unlock it first.")
        pages = [page.get_text("text") for page in doc]
    finally:
        doc.close()
    text = "\n\n".join(pages).strip()
    if not text:
        raise ValueError("No selectable text found. This PDF looks like a scan "
                         "(images only) and needs OCR.")
    with atomic_output(dst) as tmp:
        tmp.write_text(text, encoding="utf-8-sig")
    return Result([dst], note=f"Text from {len(pages)} page{'s' if len(pages) != 1 else ''}")


def pdf_to_images(src: Path, dst: Path, target: str, quality: int, dpi: int = 150) -> Result:
    from PIL import Image
    pymupdf = _pymupdf()
    doc = pymupdf.open(str(src))
    outputs = []
    try:
        if doc.is_encrypted:
            raise ValueError("This PDF is password-protected. Unlock it first.")
        count = doc.page_count
        if count == 0:
            raise ValueError("This PDF has no pages.")
        for index, page in enumerate(doc):
            pix = page.get_pixmap(dpi=dpi, alpha=False)
            img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            out = dst if count == 1 else dst.with_name(
                f"{dst.stem}_page_{index + 1:03d}{dst.suffix}")
            with atomic_output(out) as tmp:
                if target == "JPG":
                    img.save(tmp, format="JPEG", quality=quality, optimize=True)
                else:
                    img.save(tmp, format="PNG", optimize=True)
            outputs.append(out)
    finally:
        doc.close()
    note = "" if len(outputs) == 1 else f"{len(outputs)} pages saved as separate images"
    return Result(outputs, note=note)


def read_table(src: Path):
    import pandas as pd
    fmt = format_of(src.suffix)
    if fmt == "CSV":
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                df = pd.read_csv(src, sep=None, engine="python", encoding=enc)
                break
            except UnicodeDecodeError:
                continue
    elif fmt == "XLSX":
        df = pd.read_excel(src, sheet_name=0)
    elif fmt == "JSON":
        with open(src, "r", encoding="utf-8-sig") as fh:
            data = json.load(fh)
        if isinstance(data, dict):
            tables = [v for v in data.values()
                      if isinstance(v, list) and v and isinstance(v[0], dict)]
            data = tables[0] if tables else [data]
        if not isinstance(data, list):
            data = [data]
        if data and not isinstance(data[0], dict):
            df = pd.DataFrame({"value": data})
        else:
            df = pd.json_normalize(data)
    else:
        raise ValueError("Unsupported data file.")
    if df.empty and len(df.columns) == 0:
        raise ValueError("The file contains no data.")
    return df


def convert_data(src: Path, dst: Path, target: str) -> Result:
    import pandas as pd
    df = read_table(src)
    with atomic_output(dst) as tmp:
        if target == "CSV":
            df.to_csv(tmp, index=False, encoding="utf-8-sig")
        elif target == "JSON":
            df.to_json(tmp, orient="records", indent=2, force_ascii=False,
                       date_format="iso")
        elif target == "XLSX":
            from openpyxl.styles import Font
            from openpyxl.utils import get_column_letter
            with pd.ExcelWriter(tmp, engine="openpyxl") as writer:
                df.to_excel(writer, index=False, sheet_name="Sheet1")
                sheet = writer.sheets["Sheet1"]
                for idx, col in enumerate(df.columns, start=1):
                    longest = max([len(str(col))] +
                                  [len(str(v)) for v in df[col].head(500)])
                    sheet.column_dimensions[get_column_letter(idx)].width = \
                        min(max(longest + 2, 8), 60)
                    sheet.cell(row=1, column=idx).font = Font(bold=True)
        else:
            raise ValueError(f"Unsupported target: {target}")
    return Result([dst], note=f"{len(df):,} rows, {len(df.columns)} columns")


def convert(src: Path, dst: Path, target: str, quality: int = 85) -> Result:
    """Dispatch to the right converter."""
    if target not in available_targets(src.suffix):
        raise ValueError(f"{src.suffix.upper()} files cannot be converted to {target}.")
    src_fmt = format_of(src.suffix)
    if src_fmt in IMAGE_FORMATS:
        return convert_image(src, dst, target, quality)
    if src_fmt == "PDF":
        if target == "TXT":
            return pdf_to_txt(src, dst)
        return pdf_to_images(src, dst, target, quality)
    if src_fmt == "TXT":
        return txt_to_pdf(src, dst)
    return convert_data(src, dst, target)


def compress_image(src: Path, dst: Path, quality: int, scale: float) -> Result:
    from PIL import Image, ImageOps
    fmt = format_of(dst.suffix)
    with Image.open(src) as opened:
        im = ImageOps.exif_transpose(opened)
        if scale < 1:
            im = im.resize((max(1, round(im.width * scale)),
                            max(1, round(im.height * scale))), Image.LANCZOS)
        with atomic_output(dst) as tmp:
            if fmt == "JPG":
                _flatten(im).save(tmp, format="JPEG", quality=quality,
                                  optimize=True, progressive=True)
            elif fmt == "WEBP":
                im.save(tmp, format="WEBP", quality=quality, method=6)
            elif fmt == "PNG":
                out = im
                if quality < 90:    # colour reduction gives real PNG savings
                    colours = 256 if quality >= 60 else 128
                    out = im.convert("RGBA").quantize(
                        colors=colours, method=Image.Quantize.FASTOCTREE)
                out.save(tmp, format="PNG", optimize=True)
            else:
                raise ValueError("Image compression supports JPG, PNG and WEBP.")
    return Result([dst])


def optimize_pdf(src: Path, dst: Path, level: str) -> Result:
    pymupdf = _pymupdf()
    doc = pymupdf.open(str(src))
    try:
        if doc.is_encrypted:
            raise ValueError("This PDF is password-protected. Unlock it first.")
        note = ""
        if level == "Strong":
            if hasattr(doc, "rewrite_images"):
                doc.rewrite_images(dpi_threshold=150, dpi_target=110, quality=65)
            else:
                note = "Update PyMuPDF to enable image downsampling."
        with atomic_output(dst) as tmp:
            doc.save(str(tmp), garbage=4, deflate=True, deflate_images=True,
                     deflate_fonts=True, clean=True)
    finally:
        doc.close()
    return Result([dst], note=note)


def create_zip(paths: list, dst: Path, level: int) -> Result:
    used = set()
    with atomic_output(dst) as tmp:
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED, compresslevel=level) as zf:
            for path in paths:
                name, n = path.name, 1
                while name.lower() in used:
                    name = f"{path.stem} ({n}){path.suffix}"
                    n += 1
                used.add(name.lower())
                zf.write(path, arcname=name)
    return Result([dst], note=f"{len(paths)} file{'s' if len(paths) != 1 else ''} archived")


# =============================================================================
# 5.  CUSTOM WIDGETS
# =============================================================================

def round_rect(canvas, x1, y1, x2, y2, r, **kw):
    r = max(0, min(r, (x2 - x1) / 2, (y2 - y1) / 2))
    pts = [x1 + r, y1, x2 - r, y1, x2, y1, x2, y1 + r, x2, y2 - r, x2, y2,
           x2 - r, y2, x1 + r, y2, x1, y2, x1, y2 - r, x1, y1 + r, x1, y1]
    return canvas.create_polygon(pts, smooth=True, splinesteps=16, **kw)


class Button(tk.Canvas):
    STYLES = {
        "primary": dict(fill=ACCENT, hover=ACCENT_HOVER, fg=ON_ACCENT, outline=ACCENT),
        "secondary": dict(fill=CARD_ALT, hover=HOVER, fg=TEXT, outline=BORDER),
        "danger": dict(fill=CARD_ALT, hover="#3A1D1F", fg=ERROR, outline="#4A2528"),
    }

    def __init__(self, parent, text, command, kind="primary", width=160, height=40):
        super().__init__(parent, width=width, height=height, bg=parent.cget("bg"),
                         highlightthickness=0, bd=0, cursor="hand2")
        self._cmd = command
        self._style = self.STYLES[kind]
        self._enabled = True
        s = self._style
        self._shape = round_rect(self, 1, 1, width - 1, height - 1, 10,
                                 fill=s["fill"], outline=s["outline"])
        self._text = self.create_text(width / 2, height / 2, text=text,
                                      fill=s["fg"], font=F(10, "bold"))
        self.bind("<Enter>", lambda e: self._paint(True))
        self.bind("<Leave>", lambda e: self._paint(False))
        self.bind("<Button-1>", self._click)

    def _paint(self, hover: bool):
        s = self._style
        if not self._enabled:
            self.itemconfigure(self._shape, fill="#222227", outline="#222227")
            self.itemconfigure(self._text, fill=DIM)
            return
        self.itemconfigure(self._shape, fill=s["hover"] if hover else s["fill"],
                           outline=s["hover"] if hover else s["outline"])
        self.itemconfigure(self._text, fill=s["fg"])

    def _click(self, _event):
        if self._enabled:
            self._cmd()

    def set_enabled(self, enabled: bool):
        self._enabled = enabled
        self.configure(cursor="hand2" if enabled else "arrow")
        self._paint(False)

    def set_text(self, text: str):
        self.itemconfigure(self._text, text=text)


class Segmented(tk.Frame):
    """Compact option selector (like a segmented control)."""

    def __init__(self, parent, options, default, command=None):
        super().__init__(parent, bg=CARD_ALT, highlightbackground=BORDER,
                         highlightthickness=1)
        self.value = default
        self._command = command
        self._labels = {}
        for opt in options:
            lbl = tk.Label(self, text=opt, font=F(9, "bold"), padx=11, pady=6,
                           cursor="hand2")
            lbl.pack(side="left", padx=2, pady=2)
            lbl.bind("<Button-1>", lambda e, o=opt: self.select(o))
            self._labels[opt] = lbl
        self.select(default, fire=False)

    def select(self, option, fire=True):
        self.value = option
        for name, lbl in self._labels.items():
            active = name == option
            lbl.configure(bg=ACCENT if active else CARD_ALT,
                          fg=ON_ACCENT if active else MUTED)
        if fire and self._command:
            self._command(option)


class Switch(tk.Canvas):
    def __init__(self, parent, value: bool, command):
        super().__init__(parent, width=48, height=26, bg=parent.cget("bg"),
                         highlightthickness=0, cursor="hand2")
        self.value = value
        self._command = command
        self.bind("<Button-1>", self._toggle)
        self._draw()

    def _draw(self):
        self.delete("all")
        color = ACCENT if self.value else "#33333B"
        round_rect(self, 1, 1, 47, 25, 12, fill=color, outline=color)
        cx = 36 if self.value else 12
        self.create_oval(cx - 9, 4, cx + 9, 22, fill="#EDE6DC", outline="")

    def _toggle(self, _event=None):
        self.value = not self.value
        self._draw()
        self._command(self.value)


class LabeledSlider(tk.Frame):
    def __init__(self, parent, from_, to, value, suffix="", length=260, on_release=None):
        super().__init__(parent, bg=parent.cget("bg"))
        self._suffix = suffix
        self._var = value
        self.value_label = tk.Label(self, text=f"{value}{suffix}", width=5,
                                    font=F(10, "bold"), bg=parent.cget("bg"), fg=ACCENT)
        self.scale = ttk.Scale(self, from_=from_, to=to, orient="horizontal",
                               length=length, style="Orange.Horizontal.TScale",
                               command=self._changed)
        self.scale.set(value)
        self.scale.pack(side="left")
        self.value_label.pack(side="left", padx=(10, 0))
        if on_release:
            self.scale.bind("<ButtonRelease-1>", lambda e: on_release(self.get()))

    def _changed(self, raw):
        self._var = int(round(float(raw)))
        self.value_label.configure(text=f"{self._var}{self._suffix}")

    def get(self) -> int:
        return self._var

    def set(self, value: int):
        self.scale.set(value)


class NavItem(tk.Frame):
    def __init__(self, parent, text, command):
        super().__init__(parent, bg=SIDEBAR, cursor="hand2")
        self.active = False
        self.bar = tk.Frame(self, bg=SIDEBAR, width=3)
        self.bar.pack(side="left", fill="y")
        self.label = tk.Label(self, text=text, font=F(11), bg=SIDEBAR, fg=MUTED,
                              anchor="w", padx=22, pady=12)
        self.label.pack(side="left", fill="x", expand=True)
        for widget in (self, self.label):
            widget.bind("<Button-1>", lambda e: command())
            widget.bind("<Enter>", lambda e: self._paint(True))
            widget.bind("<Leave>", lambda e: self._paint(False))

    def set_active(self, flag: bool):
        self.active = flag
        self._paint(False)

    def _paint(self, hover: bool):
        bg = NAV_ACTIVE if self.active else (CARD if hover else SIDEBAR)
        self.configure(bg=bg)
        self.label.configure(bg=bg, fg=TEXT if (self.active or hover) else MUTED,
                             font=F(11, "bold" if self.active else "normal"))
        self.bar.configure(bg=ACCENT if self.active else bg)


class FormatPicker(tk.Frame):
    """Dropdown that always lists every format; unusable ones are greyed out."""

    def __init__(self, parent, on_select):
        super().__init__(parent, bg=CARD_ALT, highlightbackground=BORDER,
                         highlightthickness=1, cursor="hand2")
        self.on_select = on_select
        self.enabled = set()
        self.value = None
        self._popup = None
        self.label = tk.Label(self, text="Select format", font=F(11), bg=CARD_ALT,
                              fg=MUTED, anchor="w", padx=14, pady=11, width=30,
                              cursor="hand2")
        self.label.pack(side="left")
        self.chevron = tk.Label(self, text="\u25BE", font=F(12), bg=CARD_ALT,
                                fg=ACCENT, padx=14, cursor="hand2")
        self.chevron.pack(side="right")
        for widget in (self, self.label, self.chevron):
            widget.bind("<Button-1>", self.toggle)
        self.winfo_toplevel().bind("<Button-1>", self._on_root_click, add="+")

    # ---- state -----------------------------------------------------------
    def set_enabled(self, formats):
        self.enabled = set(formats)
        if self.value not in self.enabled:
            self.set_value(None)

    def set_value(self, fmt):
        self.value = fmt
        if fmt:
            self.label.configure(text=f"{fmt}    {FORMATS[fmt][0]}", fg=TEXT,
                                 font=F(11, "bold"))
        else:
            self.label.configure(text="Select format", fg=MUTED, font=F(11))

    # ---- popup -----------------------------------------------------------
    def toggle(self, _event=None):
        if self._popup:
            self.close()
        else:
            self.open()

    def _on_root_click(self, event):
        if self._popup and not str(event.widget).startswith(str(self)):
            self.close()

    def close(self):
        if self._popup:
            self._popup.destroy()
            self._popup = None
        self.configure(highlightbackground=BORDER)

    def open(self):
        self.configure(highlightbackground=ACCENT)
        pop = tk.Toplevel(self)
        pop.withdraw()
        pop.overrideredirect(True)
        pop.configure(bg=ACCENT)
        inner = tk.Frame(pop, bg=CARD, padx=6, pady=6)
        inner.pack(padx=1, pady=1)
        self._popup = pop

        if not self.enabled:
            tk.Label(inner, text="Choose a file first \u2013 the formats it can be\n"
                                 "converted to will be unlocked here.",
                     font=F(9), bg=CARD, fg=WARNING, justify="left",
                     padx=10, pady=8).pack(anchor="w")

        category = None
        for name, (ext, cat, desc) in FORMATS.items():
            if cat != category:
                category = cat
                tk.Label(inner, text=cat.upper(), font=F(8, "bold"), bg=CARD, fg=DIM,
                         anchor="w", padx=10).pack(fill="x", pady=(8, 2))
            self._make_row(inner, name, ext, desc, name in self.enabled)

        pop.update_idletasks()
        width = max(self.winfo_width(), pop.winfo_reqwidth())
        height = pop.winfo_reqheight()
        x = self.winfo_rootx()
        y = self.winfo_rooty() + self.winfo_height() + 4
        if y + height > self.winfo_screenheight() - 40:
            y = max(10, self.winfo_rooty() - height - 4)
        pop.geometry(f"{width}x{height}+{x}+{y}")
        pop.deiconify()
        pop.lift()

    def _make_row(self, parent, name, ext, desc, enabled):
        row = tk.Frame(parent, bg=CARD, cursor="hand2" if enabled else "arrow")
        row.pack(fill="x")
        fg_main = TEXT if enabled else DIM
        fg_sub = MUTED if enabled else "#47454D"
        n = tk.Label(row, text=name, font=F(10, "bold"), width=6, anchor="w",
                     bg=CARD, fg=ACCENT if enabled else DIM, padx=10, pady=6)
        d = tk.Label(row, text=desc, font=F(10), anchor="w", bg=CARD, fg=fg_main)
        e = tk.Label(row, text=ext, font=F(9), bg=CARD, fg=fg_sub, padx=12)
        n.pack(side="left")
        d.pack(side="left", fill="x", expand=True)
        e.pack(side="right")
        if not enabled:
            return
        widgets = (row, n, d, e)

        def paint(bg):
            for w in widgets:
                w.configure(bg=bg)

        for w in widgets:
            w.bind("<Enter>", lambda ev: paint(ACCENT_DIM))
            w.bind("<Leave>", lambda ev: paint(CARD))
            w.bind("<Button-1>", lambda ev, f=name: self._choose(f))

    def _choose(self, fmt):
        self.close()
        self.set_value(fmt)
        self.on_select(fmt)


def card(parent, bg=CARD):
    return tk.Frame(parent, bg=bg, highlightbackground=BORDER,
                    highlightcolor=BORDER, highlightthickness=1)


def page_header(parent, title, subtitle):
    box = tk.Frame(parent, bg=BG)
    box.pack(fill="x", pady=(0, 22))
    tk.Label(box, text=title, font=F(24, "bold"), bg=BG, fg=TEXT).pack(anchor="w")
    tk.Label(box, text=subtitle, font=F(11), bg=BG, fg=MUTED).pack(anchor="w", pady=(4, 0))


def step_header(parent, step, title):
    tk.Label(parent, text=step, font=F(8, "bold"), bg=parent.cget("bg"),
             fg=ACCENT).pack(anchor="w", padx=24, pady=(20, 2))
    tk.Label(parent, text=title, font=F(14, "bold"), bg=parent.cget("bg"),
             fg=TEXT).pack(anchor="w", padx=24, pady=(0, 14))


# =============================================================================
# 6.  APPLICATION
# =============================================================================

class FileXApp(tk.Tk):
    def __init__(self):
        super().__init__()
        global FONT_FAMILY
        families = set(tkfont.families(self))
        for name in ("Segoe UI", "SF Pro Text", "Helvetica Neue", "Inter",
                     "DejaVu Sans", "Arial"):
            if name in families:
                FONT_FAMILY = name
                break

        self.settings = Settings()
        self.history = History()
        self.busy = False
        self.src = None            # file chosen on the Convert page
        self.target = None         # chosen output format
        self.last_output = None
        self.cmp_image = None
        self.cmp_pdf = None
        self.cmp_zip = []
        self.pages = {}
        self.nav = {}
        self.current_page = None

        self._setup_window()
        self._setup_styles()
        self._build_layout()
        self.show_page("Convert")

        missing = missing_dependencies()
        if missing:
            self.after(400, lambda: messagebox.showwarning(
                "Missing libraries",
                "Some features need extra packages:\n\n"
                f"    pip install {' '.join(missing)}\n\n"
                "Install them and restart FileX."))

    # ---------------------------------------------------------------- setup
    def _setup_window(self):
        self.title(f"{APP_NAME} \u2014 File Converter & Compressor")
        self.configure(bg=BG)
        self.minsize(1060, 700)
        w, h = 1220, 780
        x = max(0, (self.winfo_screenwidth() - w) // 2)
        y = max(0, (self.winfo_screenheight() - h) // 2 - 20)
        self.geometry(f"{w}x{h}+{x}+{y}")
        self.bind_all("<Control-o>", lambda e: self.browse_file()
                      if self.current_page == "Convert" else None)

    def _setup_styles(self):
        s = ttk.Style(self)
        s.theme_use("clam")
        s.layout("Treeview", [("Treeview.treearea", {"sticky": "nswe"})])
        s.configure("Treeview", background=CARD, fieldbackground=CARD,
                    foreground=TEXT, borderwidth=0, rowheight=38, font=F(10))
        s.configure("Treeview.Heading", background=CARD_ALT, foreground=MUTED,
                    relief="flat", font=F(9, "bold"), padding=(12, 9), borderwidth=0)
        s.map("Treeview", background=[("selected", ACCENT_DIM)],
              foreground=[("selected", TEXT)])
        s.map("Treeview.Heading", background=[("active", CARD_ALT)])
        s.configure("Vertical.TScrollbar", background=HOVER, troughcolor=CARD,
                    bordercolor=CARD, arrowcolor=MUTED, relief="flat", gripcount=0)
        s.map("Vertical.TScrollbar", background=[("active", DIM)])
        s.configure("Orange.Horizontal.TProgressbar", troughcolor=CARD_ALT,
                    background=ACCENT, bordercolor=SIDEBAR, lightcolor=ACCENT,
                    darkcolor=ACCENT, thickness=6)
        s.configure("Orange.Horizontal.TScale", background=ACCENT,
                    troughcolor=CARD_ALT, bordercolor=BORDER, lightcolor=ACCENT,
                    darkcolor=ACCENT, gripcount=0)
        s.map("Orange.Horizontal.TScale", background=[("active", ACCENT_HOVER)])

    def _build_layout(self):
        body = tk.Frame(self, bg=BG)
        body.pack(fill="both", expand=True)

        # ---- sidebar
        side = tk.Frame(body, bg=SIDEBAR, width=236)
        side.pack(side="left", fill="y")
        side.pack_propagate(False)
        tk.Frame(body, bg=BORDER, width=1).pack(side="left", fill="y")

        brand = tk.Frame(side, bg=SIDEBAR)
        brand.pack(fill="x", padx=24, pady=(30, 34))
        logo = tk.Canvas(brand, width=42, height=42, bg=SIDEBAR, highlightthickness=0)
        round_rect(logo, 1, 1, 41, 41, 11, fill=ACCENT, outline=ACCENT)
        logo.create_text(21, 21, text="F", font=F(18, "bold"), fill=ON_ACCENT)
        logo.pack(side="left")
        names = tk.Frame(brand, bg=SIDEBAR)
        names.pack(side="left", padx=12)
        tk.Label(names, text=APP_NAME, font=F(18, "bold"), bg=SIDEBAR, fg=TEXT).pack(anchor="w")
        tk.Label(names, text="Convert \u2022 Compress", font=F(8), bg=SIDEBAR,
                 fg=MUTED).pack(anchor="w")

        tk.Label(side, text="WORKSPACE", font=F(8, "bold"), bg=SIDEBAR, fg=DIM,
                 anchor="w").pack(fill="x", padx=26, pady=(0, 6))
        for name in ("Convert", "Compress", "History", "Settings", "Help"):
            item = NavItem(side, name, lambda n=name: self.show_page(n))
            item.pack(fill="x")
            self.nav[name] = item

        tk.Label(side, text=f"Version {APP_VERSION}", font=F(9), bg=SIDEBAR,
                 fg=DIM).pack(side="bottom", anchor="w", padx=26, pady=22)

        # ---- main area
        main = tk.Frame(body, bg=BG)
        main.pack(side="left", fill="both", expand=True)

        status = tk.Frame(main, bg=SIDEBAR, height=38)
        status.pack(side="bottom", fill="x")
        status.pack_propagate(False)
        tk.Frame(main, bg=BORDER, height=1).pack(side="bottom", fill="x")
        self.status_dot = tk.Label(status, text="\u25CF", font=F(9), bg=SIDEBAR, fg=MUTED)
        self.status_dot.pack(side="left", padx=(24, 8))
        self.status_text = tk.Label(status, text="Ready", font=F(10), bg=SIDEBAR, fg=MUTED)
        self.status_text.pack(side="left")
        self.progress = ttk.Progressbar(status, mode="indeterminate", length=160,
                                        style="Orange.Horizontal.TProgressbar")

        self.content = tk.Frame(main, bg=BG)
        self.content.pack(fill="both", expand=True, padx=44, pady=(30, 20))

        self.pages["Convert"] = self._build_convert(self.content)
        self.pages["Compress"] = self._build_compress(self.content)
        self.pages["History"] = self._build_history(self.content)
        self.pages["Settings"] = self._build_settings(self.content)
        self.pages["Help"] = self._build_help(self.content)

    # ----------------------------------------------------------- navigation
    def show_page(self, name):
        if self.current_page:
            self.pages[self.current_page].pack_forget()
        self.picker.close()
        self.pages[name].pack(fill="both", expand=True)
        for key, item in self.nav.items():
            item.set_active(key == name)
        self.current_page = name
        if name == "History":
            self.refresh_history()

    # --------------------------------------------------------------- status
    def set_status(self, text, kind="info"):
        colors = {"info": MUTED, "ok": SUCCESS, "err": ERROR, "busy": ACCENT}
        self.status_dot.configure(fg=colors[kind])
        self.status_text.configure(text=text, fg=TEXT if kind != "info" else MUTED)

    def set_busy(self, flag, text=""):
        self.busy = flag
        if flag:
            self.progress.pack(side="right", padx=24)
            self.progress.start(12)
            self.set_status(text, "busy")
        else:
            self.progress.stop()
            self.progress.pack_forget()

    def run_task(self, work, on_success, on_error, busy_text):
        """Run work() in a background thread so the window never freezes."""
        self.set_busy(True, busy_text)
        results = queue.Queue()

        def target():
            try:
                results.put(("ok", work()))
            except Exception as exc:                       # noqa: BLE001
                results.put(("err", exc))

        threading.Thread(target=target, daemon=True).start()

        def poll():
            try:
                kind, payload = results.get_nowait()
            except queue.Empty:
                self.after(80, poll)
                return
            self.set_busy(False)
            (on_success if kind == "ok" else on_error)(payload)

        poll()

    def fail(self, title):
        def handler(exc):
            self.set_status(f"{title}.", "err")
            messagebox.showerror(title, friendly_error(exc))
        return handler

    # ------------------------------------------------------------ output path
    def resolve_output(self, src: Path, name: str, ext: str, title: str):
        folder = Path(self.settings["output_folder"])
        if self.settings["ask_location"]:
            chosen = filedialog.asksaveasfilename(
                title=title,
                initialdir=str(folder if folder.is_dir() else src.parent),
                initialfile=name, defaultextension=ext,
                filetypes=[(f"{ext[1:].upper()} file", f"*{ext}")])
            if not chosen:
                return None
            out = Path(chosen)
            if out.suffix.lower() != ext:
                out = out.with_suffix(ext)
        else:
            try:
                folder.mkdir(parents=True, exist_ok=True)
            except OSError as exc:
                messagebox.showerror("Output folder", friendly_error(exc))
                return None
            out = unique_path(folder / name)
        try:
            if out.resolve() == src.resolve():
                messagebox.showerror("Invalid destination",
                                     "The output cannot overwrite the original file.")
                return None
        except OSError:
            pass
        return out

    # =========================================================== CONVERT PAGE
    def _build_convert(self, parent):
        page = tk.Frame(parent, bg=BG)
        page_header(page, "Convert files",
                    "Switch between image, document and data formats in a few clicks.")

        # Step 1
        c1 = card(page)
        c1.pack(fill="x", pady=(0, 16))
        step_header(c1, "STEP 1", "Choose your file")
        row = tk.Frame(c1, bg=CARD)
        row.pack(fill="x", padx=24, pady=(0, 22))
        self.badge = tk.Canvas(row, width=60, height=60, bg=CARD, highlightthickness=0)
        round_rect(self.badge, 1, 1, 59, 59, 12, fill=CARD_ALT, outline=BORDER)
        self.badge_text = self.badge.create_text(30, 30, text="\u2014", fill=MUTED,
                                                 font=F(11, "bold"))
        self.badge.pack(side="left")
        info = tk.Frame(row, bg=CARD)
        info.pack(side="left", fill="x", expand=True, padx=16)
        self.file_name = tk.Label(info, text="No file selected", font=F(13, "bold"),
                                  bg=CARD, fg=TEXT, anchor="w")
        self.file_name.pack(fill="x")
        self.file_meta = tk.Label(info, text="Images, PDF, TXT, CSV, XLSX and JSON are supported",
                                  font=F(10), bg=CARD, fg=MUTED, anchor="w")
        self.file_meta.pack(fill="x", pady=(3, 0))
        Button(row, "Browse files", self.browse_file, width=150, height=42).pack(side="right")

        # Step 2
        c2 = card(page)
        c2.pack(fill="x", pady=(0, 18))
        step_header(c2, "STEP 2", "Choose output format")
        body = tk.Frame(c2, bg=CARD)
        body.pack(fill="x", padx=24, pady=(0, 22))
        self.picker = FormatPicker(body, self.on_format_selected)
        self.picker.pack(anchor="w")
        self.format_hint = tk.Label(body, text="Select a file to see which formats are available.",
                                    font=F(9), bg=CARD, fg=MUTED, anchor="w")
        self.format_hint.pack(anchor="w", pady=(8, 0))
        self.quality_holder = tk.Frame(body, bg=CARD)
        self.quality_holder.pack(fill="x")
        self.quality_row = tk.Frame(self.quality_holder, bg=CARD)
        tk.Label(self.quality_row, text="Output quality", font=F(10, "bold"),
                 bg=CARD, fg=TEXT).pack(anchor="w", pady=(18, 4))
        self.convert_quality = LabeledSlider(self.quality_row, 20, 100,
                                             int(self.settings["quality"]), suffix="%")
        self.convert_quality.pack(anchor="w")

        # Actions
        actions = tk.Frame(page, bg=BG)
        actions.pack(fill="x")
        self.convert_btn = Button(actions, "Convert file  \u2192", self.do_convert,
                                  width=210, height=48)
        self.convert_btn.pack(side="left")
        Button(actions, "Reset", self.reset_convert, kind="secondary",
               width=110, height=48).pack(side="left", padx=12)

        # Result card
        self.result_card = card(page)
        rc = tk.Frame(self.result_card, bg=CARD)
        rc.pack(fill="x", padx=24, pady=20)
        self.result_title = tk.Label(rc, text="", font=F(12, "bold"), bg=CARD, fg=SUCCESS)
        self.result_title.pack(anchor="w")
        self.result_file = tk.Label(rc, text="", font=F(11, "bold"), bg=CARD, fg=TEXT)
        self.result_file.pack(anchor="w", pady=(6, 0))
        self.result_sizes = tk.Label(rc, text="", font=F(10), bg=CARD, fg=MUTED)
        self.result_sizes.pack(anchor="w", pady=(2, 0))
        self.result_note = tk.Label(rc, text="", font=F(9), bg=CARD, fg=WARNING)
        self.result_note.pack(anchor="w")
        btns = tk.Frame(rc, bg=CARD)
        btns.pack(anchor="w", pady=(14, 0))
        Button(btns, "Open file", lambda: self._open_last(False),
               width=130, height=38).pack(side="left")
        Button(btns, "Show in folder", lambda: self._open_last(True), kind="secondary",
               width=150, height=38).pack(side="left", padx=10)
        return page

    def browse_file(self):
        patterns = " ".join(f"*{e}" for e in INPUT_EXTENSIONS)
        path = filedialog.askopenfilename(
            title="Select a file",
            filetypes=[("Supported files", patterns), ("All files", "*.*")])
        if path:
            self.set_source(Path(path))

    def set_source(self, path: Path):
        ext = path.suffix.lower()
        if not path.is_file():
            messagebox.showerror("File not found", "That file does not exist.")
            return
        if not available_targets(ext):
            messagebox.showerror("Unsupported file",
                                 f"{ext or 'This file type'} is not supported by FileX.")
            return
        if path.stat().st_size == 0:
            messagebox.showwarning("Empty file", "The selected file is empty.")
            return
        self.src = path
        self.target = None
        self.badge.itemconfigure(self.badge_text, text=ext[1:].upper()[:4], fill=ACCENT)
        self.file_name.configure(text=shorten(path.name, 56))
        self.file_meta.configure(
            text=f"{FORMATS[format_of(ext)][2]}  \u2022  {human_size(path.stat().st_size)}")
        targets = available_targets(ext)
        self.picker.set_enabled(targets)
        self.format_hint.configure(
            text="Available for this file:  " + ",  ".join(targets), fg=MUTED)
        self.quality_row.pack_forget()
        self.hide_result()
        self.set_status(f"Loaded {path.name}", "ok")

    def reset_convert(self):
        self.src = None
        self.target = None
        self.badge.itemconfigure(self.badge_text, text="\u2014", fill=MUTED)
        self.file_name.configure(text="No file selected")
        self.file_meta.configure(text="Images, PDF, TXT, CSV, XLSX and JSON are supported")
        self.picker.set_enabled([])
        self.format_hint.configure(text="Select a file to see which formats are available.")
        self.quality_row.pack_forget()
        self.hide_result()
        self.set_status("Ready")

    def on_format_selected(self, fmt):
        self.target = fmt
        self.hide_result()
        if fmt in ("JPG", "WEBP"):
            self.quality_row.pack(fill="x")
        else:
            self.quality_row.pack_forget()
        self.set_status(f"Output format: {fmt}")

    def do_convert(self):
        if self.busy:
            return
        if not self.src:
            messagebox.showwarning("No file selected", "Please choose a file first.")
            return
        if not self.target:
            messagebox.showwarning("Choose a format",
                                   "Open the format menu and pick an output format.")
            return
        src, target = self.src, self.target
        if not src.is_file():
            messagebox.showerror("File not found", "The selected file no longer exists.")
            return
        ext = FORMATS[target][0]
        dst = self.resolve_output(src, f"{src.stem}{ext}", ext, f"Save {target} file")
        if dst is None:
            return
        quality = self.convert_quality.get()
        self.hide_result()
        self.run_task(
            lambda: convert(src, dst, target, quality),
            lambda res: self.convert_done(src, target, res),
            self.fail("Conversion failed"),
            f"Converting {src.name} to {target}\u2026")

    def convert_done(self, src, target, res):
        self.history.add(f"Convert to {target}", src, res.primary)
        self.show_result("Conversion complete", res, src.stat().st_size)
        self.set_status(f"Saved {res.primary.name}", "ok")

    def show_result(self, title, res, original_size):
        self.last_output = res.primary
        out_size = res.total_size
        self.result_title.configure(text="\u2713  " + title)
        self.result_file.configure(text=shorten(res.primary.name, 60))
        self.result_sizes.configure(
            text=f"{human_size(original_size)}  \u2192  {human_size(out_size)}"
                 f"   ({size_change(original_size, out_size)})   \u2022   "
                 f"{shorten(str(res.primary.parent), 60)}")
        self.result_note.configure(text=res.note)
        self.result_card.pack(fill="x", pady=(18, 0))

    def hide_result(self):
        self.result_card.pack_forget()

    def _open_last(self, folder: bool):
        if not self.last_output or not Path(self.last_output).exists():
            messagebox.showwarning("Not found", "The output file no longer exists.")
            return
        try:
            (reveal_path if folder else open_path)(self.last_output)
        except OSError as exc:
            messagebox.showerror("Could not open", friendly_error(exc))

    # =========================================================== COMPRESS PAGE
    def _build_compress(self, parent):
        page = tk.Frame(parent, bg=BG)
        page_header(page, "Compress files",
                    "Make images, PDFs and folders of files smaller without the fuss.")
        grid = tk.Frame(page, bg=BG)
        grid.pack(fill="both", expand=True)
        for i in range(3):
            grid.columnconfigure(i, weight=1, uniform="tools")
        grid.rowconfigure(0, weight=1)

        # ----- image
        body = self._tool_card(grid, 0, "Compress image",
                               "Reduce JPG, PNG and WEBP size with minimal visible loss.")
        self.img_run = Button(body, "Compress image", self.run_compress_image,
                              width=226, height=44)
        self.img_run.pack(side="bottom", anchor="w")
        self.img_result = self._result_label(body)
        self.img_file = self._file_picker(body, "Choose image", self.pick_image)
        tk.Label(body, text="Quality", font=F(10, "bold"), bg=CARD,
                 fg=TEXT).pack(anchor="w", pady=(16, 4))
        self.img_quality = LabeledSlider(body, 20, 95, 75, suffix="%", length=170)
        self.img_quality.pack(anchor="w")
        tk.Label(body, text="Resize", font=F(10, "bold"), bg=CARD,
                 fg=TEXT).pack(anchor="w", pady=(14, 6))
        self.img_scale = Segmented(body, ["100%", "75%", "50%", "25%"], "100%")
        self.img_scale.pack(anchor="w")

        # ----- pdf
        body = self._tool_card(grid, 1, "Optimize PDF",
                               "Clean up structure and compress streams to shrink PDFs.")
        Button(body, "Optimize PDF", self.run_optimize_pdf, width=226,
               height=44).pack(side="bottom", anchor="w")
        self.pdf_result = self._result_label(body)
        self.pdf_file = self._file_picker(body, "Choose PDF", self.pick_pdf)
        tk.Label(body, text="Level", font=F(10, "bold"), bg=CARD,
                 fg=TEXT).pack(anchor="w", pady=(16, 6))
        self.pdf_level = Segmented(body, ["Standard", "Strong"], "Standard")
        self.pdf_level.pack(anchor="w")
        tk.Label(body, text="Standard: lossless clean-up.\nStrong: also downsamples\n"
                            "large images for bigger savings.",
                 font=F(9), bg=CARD, fg=MUTED, justify="left").pack(anchor="w", pady=(10, 0))

        # ----- zip
        body = self._tool_card(grid, 2, "Create ZIP",
                               "Bundle several files into a single compressed archive.")
        Button(body, "Create ZIP", self.run_zip, width=226,
               height=44).pack(side="bottom", anchor="w")
        self.zip_result = self._result_label(body)
        self.zip_files = self._file_picker(body, "Choose files", self.pick_zip_files,
                                           empty="No files selected")
        tk.Label(body, text="Compression", font=F(10, "bold"), bg=CARD,
                 fg=TEXT).pack(anchor="w", pady=(16, 6))
        self.zip_level = Segmented(body, ["Fast", "Normal", "Max"], "Normal")
        self.zip_level.pack(anchor="w")
        return page

    def _tool_card(self, parent, col, title, desc):
        c = card(parent)
        c.grid(row=0, column=col, sticky="nsew",
               padx=(0 if col == 0 else 8, 0 if col == 2 else 8))
        tk.Label(c, text=title, font=F(14, "bold"), bg=CARD, fg=TEXT).pack(
            anchor="w", padx=22, pady=(22, 4))
        tk.Label(c, text=desc, font=F(9), bg=CARD, fg=MUTED, justify="left",
                 wraplength=236).pack(anchor="w", padx=22, pady=(0, 16))
        body = tk.Frame(c, bg=CARD)
        body.pack(fill="both", expand=True, padx=22, pady=(0, 22))
        return body

    def _file_picker(self, body, button_text, command, empty="No file selected"):
        label = tk.Label(body, text=empty, font=F(10), bg=CARD_ALT, fg=MUTED,
                         anchor="w", padx=12, pady=10, width=27)
        label.pack(anchor="w", pady=(0, 8))
        Button(body, button_text, command, kind="secondary", width=226,
               height=38).pack(anchor="w")
        return label

    def _result_label(self, body):
        label = tk.Label(body, text="", font=F(9), bg=CARD, fg=MUTED, justify="left",
                         wraplength=226, anchor="w")
        label.pack(side="bottom", anchor="w", pady=(0, 12))
        return label

    # ---- compress: image
    def pick_image(self):
        path = filedialog.askopenfilename(
            title="Choose an image",
            filetypes=[("Images", "*.jpg *.jpeg *.png *.webp")])
        if path:
            self.cmp_image = Path(path)
            self.img_file.configure(text=shorten(self.cmp_image.name, 28), fg=TEXT)
            self.img_result.configure(text="")

    def run_compress_image(self):
        if self.busy:
            return
        src = self.cmp_image
        if not src or not src.is_file():
            messagebox.showwarning("Select an image", "Please choose a JPG, PNG or WEBP image.")
            return
        ext = ".jpg" if src.suffix.lower() in (".jpg", ".jpeg") else src.suffix.lower()
        if ext not in (".jpg", ".png", ".webp"):
            messagebox.showwarning("Unsupported", "Image compression supports JPG, PNG and WEBP.")
            return
        dst = self.resolve_output(src, f"{src.stem}_compressed{ext}", ext, "Save compressed image")
        if dst is None:
            return
        quality = self.img_quality.get()
        scale = int(self.img_scale.value.rstrip("%")) / 100
        self.run_task(
            lambda: compress_image(src, dst, quality, scale),
            lambda res: self.tool_done(self.img_result, "Image compression", src, res),
            self.fail("Compression failed"), f"Compressing {src.name}\u2026")

    # ---- compress: pdf
    def pick_pdf(self):
        path = filedialog.askopenfilename(title="Choose a PDF",
                                          filetypes=[("PDF files", "*.pdf")])
        if path:
            self.cmp_pdf = Path(path)
            self.pdf_file.configure(text=shorten(self.cmp_pdf.name, 28), fg=TEXT)
            self.pdf_result.configure(text="")

    def run_optimize_pdf(self):
        if self.busy:
            return
        src = self.cmp_pdf
        if not src or not src.is_file():
            messagebox.showwarning("Select a PDF", "Please choose a PDF file first.")
            return
        dst = self.resolve_output(src, f"{src.stem}_optimized.pdf", ".pdf", "Save optimized PDF")
        if dst is None:
            return
        level = self.pdf_level.value
        self.run_task(
            lambda: optimize_pdf(src, dst, level),
            lambda res: self.tool_done(self.pdf_result, "PDF optimization", src, res),
            self.fail("PDF optimization failed"), f"Optimizing {src.name}\u2026")

    # ---- compress: zip
    def pick_zip_files(self):
        paths = filedialog.askopenfilenames(title="Select files to archive")
        if paths:
            self.cmp_zip = [Path(p) for p in paths]
            total = sum(p.stat().st_size for p in self.cmp_zip if p.is_file())
            self.zip_files.configure(
                text=f"{len(self.cmp_zip)} files \u2022 {human_size(total)}", fg=TEXT)
            self.zip_result.configure(text="")

    def run_zip(self):
        if self.busy:
            return
        files = [p for p in self.cmp_zip if p.is_file()]
        if not files:
            messagebox.showwarning("No files", "Please choose at least one file.")
            return
        dst = self.resolve_output(files[0], "FileX_archive.zip", ".zip", "Save ZIP archive")
        if dst is None:
            return
        if any(dst.resolve() == p.resolve() for p in files):
            messagebox.showerror("Invalid destination",
                                 "The ZIP cannot overwrite one of the selected files.")
            return
        level = {"Fast": 1, "Normal": 6, "Max": 9}[self.zip_level.value]
        total = sum(p.stat().st_size for p in files)
        self.run_task(
            lambda: create_zip(files, dst, level),
            lambda res: self.tool_done(self.zip_result, "ZIP creation", files[0], res, total),
            self.fail("ZIP creation failed"), "Creating ZIP archive\u2026")

    def tool_done(self, label, operation, src, res, original=None):
        original = original if original is not None else src.stat().st_size
        out_size = res.total_size
        self.history.add(operation, src, res.primary)
        change = size_change(original, out_size)
        if out_size < original:
            text, color = (f"\u2713 {human_size(original)} \u2192 {human_size(out_size)}  ({change})", SUCCESS)
        else:
            text, color = (f"{human_size(original)} \u2192 {human_size(out_size)}  ({change}). "
                           "No saving \u2013 try a lower setting.", WARNING)
        if res.note:
            text += f"\n{res.note}"
        label.configure(text=text, fg=color)
        self.last_output = res.primary
        self.set_status(f"Saved {res.primary.name}", "ok")

    # ============================================================ HISTORY PAGE
    def _build_history(self, parent):
        page = tk.Frame(parent, bg=BG)
        page_header(page, "History", "Your latest conversions and compression tasks.")

        bar = tk.Frame(page, bg=BG)
        bar.pack(side="bottom", fill="x", pady=(14, 0))
        Button(bar, "Open file", lambda: self._history_open(False), width=130,
               height=40).pack(side="left")
        Button(bar, "Show in folder", lambda: self._history_open(True), kind="secondary",
               width=150, height=40).pack(side="left", padx=10)
        Button(bar, "Clear history", self.clear_history, kind="danger", width=140,
               height=40).pack(side="right")

        wrap = card(page)
        wrap.pack(fill="both", expand=True)
        cols = ("time", "operation", "input", "output")
        self.tree = ttk.Treeview(wrap, columns=cols, show="headings", selectmode="browse")
        for col, title, width in (("time", "Date / time", 170), ("operation", "Operation", 170),
                                  ("input", "Input file", 220), ("output", "Output", 400)):
            self.tree.heading(col, text=title, anchor="w")
            self.tree.column(col, width=width, anchor="w")
        scroll = ttk.Scrollbar(wrap, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=scroll.set)
        scroll.pack(side="right", fill="y")
        self.tree.pack(fill="both", expand=True)
        self.tree.bind("<Double-1>", lambda e: self._history_open(False))
        self.history_empty = tk.Label(wrap, text="No activity yet.\nYour conversions will appear here.",
                                      font=F(11), bg=CARD, fg=DIM, justify="center")
        return page

    def refresh_history(self):
        self.tree.delete(*self.tree.get_children())
        for i, entry in enumerate(self.history.items):
            self.tree.insert("", "end", iid=str(i), values=(
                entry.get("time", ""), entry.get("operation", ""),
                entry.get("input", ""), entry.get("output", "")))
        if self.history.items:
            self.history_empty.place_forget()
        else:
            self.history_empty.place(relx=0.5, rely=0.5, anchor="center")

    def _history_open(self, folder: bool):
        selected = self.tree.selection()
        if not selected:
            messagebox.showinfo("History", "Select an entry first.")
            return
        path = Path(self.history.items[int(selected[0])]["output"])
        if not path.exists():
            messagebox.showwarning("Not found", "This file has been moved or deleted.")
            return
        try:
            (reveal_path if folder else open_path)(path)
        except OSError as exc:
            messagebox.showerror("Could not open", friendly_error(exc))

    def clear_history(self):
        if self.history.items and messagebox.askyesno(
                "Clear history", "Remove all history entries?\nYour files are not deleted."):
            self.history.clear()
            self.refresh_history()

    # ============================================================ SETTINGS PAGE
    def _build_settings(self, parent):
        page = tk.Frame(parent, bg=BG)
        page_header(page, "Settings", "Make FileX work the way you do.")
        wrap = card(page)
        wrap.pack(fill="x")

        def row(title, desc, last=False):
            r = tk.Frame(wrap, bg=CARD)
            r.pack(fill="x", padx=24, pady=18)
            left = tk.Frame(r, bg=CARD)
            left.pack(side="left", fill="x", expand=True)
            tk.Label(left, text=title, font=F(11, "bold"), bg=CARD, fg=TEXT).pack(anchor="w")
            desc_label = tk.Label(left, text=desc, font=F(9), bg=CARD, fg=MUTED,
                                  justify="left", anchor="w")
            desc_label.pack(anchor="w", pady=(2, 0))
            if not last:
                tk.Frame(wrap, bg=BORDER, height=1).pack(fill="x", padx=24)
            return r, desc_label

        r, self.folder_desc = row("Output folder", self.settings["output_folder"])
        Button(r, "Change", self.change_folder, kind="secondary", width=120,
               height=38).pack(side="right")

        r, _ = row("Ask where to save each file",
                   "When off, results are saved automatically to the output folder.")
        Switch(r, bool(self.settings["ask_location"]), self._set_ask).pack(side="right")

        r, _ = row("Default quality",
                   "Starting quality for JPG / WEBP conversions.", last=True)
        LabeledSlider(r, 20, 100, int(self.settings["quality"]), suffix="%", length=200,
                      on_release=self._set_quality).pack(side="right")

        actions = tk.Frame(page, bg=BG)
        actions.pack(fill="x", pady=(18, 0))
        Button(actions, "Open output folder", self.open_output_folder, kind="secondary",
               width=190, height=40).pack(side="left")
        return page

    def change_folder(self):
        folder = filedialog.askdirectory(title="Choose output folder",
                                         initialdir=self.settings["output_folder"])
        if folder:
            self.settings["output_folder"] = os.path.abspath(folder)
            self.settings.save()
            self.folder_desc.configure(text=self.settings["output_folder"])
            self.set_status("Output folder updated", "ok")

    def _set_ask(self, value):
        self.settings["ask_location"] = value
        self.settings.save()

    def _set_quality(self, value):
        self.settings["quality"] = value
        self.settings.save()
        self.convert_quality.set(value)

    def open_output_folder(self):
        folder = Path(self.settings["output_folder"])
        try:
            folder.mkdir(parents=True, exist_ok=True)
            open_path(folder)
        except OSError as exc:
            messagebox.showerror("Could not open folder", friendly_error(exc))

    # ================================================================ HELP PAGE
    def _build_help(self, parent):
        page = tk.Frame(parent, bg=BG)
        page_header(page, "Help", "What FileX can do and how to use it.")
        wrap = card(page)
        wrap.pack(fill="x")
        tk.Label(wrap, text="SUPPORTED CONVERSIONS", font=F(8, "bold"), bg=CARD,
                 fg=ACCENT).pack(anchor="w", padx=24, pady=(20, 10))
        rows = [
            ("Images  (JPG, PNG, WEBP, BMP, GIF, TIFF, ICO)", "Any image format, or PDF"),
            ("PDF", "TXT text, or PNG / JPG images (one per page)"),
            ("TXT", "PDF"),
            ("CSV", "XLSX, JSON"),
            ("XLSX", "CSV, JSON"),
            ("JSON", "CSV, XLSX"),
        ]
        for src, dst in rows:
            r = tk.Frame(wrap, bg=CARD)
            r.pack(fill="x", padx=24, pady=5)
            tk.Label(r, text=src, font=F(10, "bold"), bg=CARD, fg=TEXT, width=44,
                     anchor="w").pack(side="left")
            tk.Label(r, text="\u2192", font=F(10), bg=CARD, fg=ACCENT).pack(side="left", padx=10)
            tk.Label(r, text=dst, font=F(10), bg=CARD, fg=MUTED, anchor="w").pack(side="left")
        tk.Label(wrap, text="QUICK TIPS", font=F(8, "bold"), bg=CARD,
                 fg=ACCENT).pack(anchor="w", padx=24, pady=(22, 8))
        tips = ("1.  Click Browse files (or press Ctrl+O) and choose your file.\n"
                "2.  Open the format menu \u2013 formats not possible for your file are greyed out.\n"
                "3.  Press Convert. The result is saved to your output folder (see Settings).\n"
                "4.  Scanned PDFs have no text, so PDF \u2192 TXT needs an OCR tool for those.\n"
                "5.  Original files are never changed or overwritten.")
        tk.Label(wrap, text=tips, font=F(10), bg=CARD, fg=MUTED, justify="left").pack(
            anchor="w", padx=24, pady=(0, 24))
        return page


# =============================================================================
# 7.  ENTRY POINT
# =============================================================================

def main():
    if sys.platform == "win32":
        try:                                    # crisp text on high-DPI screens
            import ctypes
            ctypes.windll.shcore.SetProcessDpiAwareness(1)
        except Exception:                       # noqa: BLE001
            pass
    FileXApp().mainloop()


if __name__ == "__main__":
    main()